from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    # Create presentation
    prs = Presentation()
    
    # Define colors
    primary_color = RGBColor(30, 86, 49)      # Deep Green
    text_color = RGBColor(33, 33, 33)         # Dark Gray
    
    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "ExplainableCropAI"
    subtitle.text = "Interpretable Deep Learning for Agricultural Disease Detection\n\nSantosh Kumar, Durgesh Nandan, Sumit Kumar\nSymbiosis Institute of Technology\nMay 2025"
    
    # Slide 2: Problem
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "The Problem: Black-Box AI in Agriculture"
    content.text = """• Challenge: AI systems function as "black boxes"
• Impact: Limited farmer trust and adoption
• Key Question: How to make AI transparent and trustworthy?"""
    
    # Slide 3: Solution
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Our Solution: ExplainableCropAI"
    content.text = """• Multi-level Explanation System
• 90% Disease Classification Accuracy
• Visual heatmaps + Natural language + Uncertainty
• Context-aware recommendations"""
    
    # Slide 4: Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "System Architecture"
    content.text = """1. Disease Detection Model (Modified ResNet50)
2. Visual Explanation Generator (Grad-CAM)
3. Uncertainty Quantification (Monte Carlo Dropout)
4. Natural Language Generator
5. Recommendation Engine"""
    
    # Slide 5: Results
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Technical Results"
    content.text = """Disease Detection Performance:
• Accuracy: 90.0%
• Precision: 89.5%
• Recall: 90.2%
• F1-Score: 89.8%

Explanation Quality:
• Effective visual explanations
• Well-calibrated uncertainty
• Positive user evaluation"""
    
    # Slide 6: User Evaluation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "User Evaluation Results"
    content.text = """• Decision Support: Improved accuracy across user groups
• Trust Calibration: Increased trust with explanations
• Cognitive Load: Reduced mental effort
• Learning Value: Educational benefits beyond decisions"""
    
    # Slide 7: Contributions
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Key Contributions"
    content.text = """• Novel multi-level explanation framework
• Domain-specific knowledge integration
• Empirical validation of explainability impact
• Blueprint for responsible AI in agriculture"""
    
    # Slide 8: Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Conclusion"
    content.text = """• ExplainableCropAI bridges AI accuracy and interpretability
• 90% accuracy with comprehensive explanations
• Improved trust and decision-making
• Foundation for responsible AI deployment

"Transforming AI from black-box to assistive technology" """
    
    # Slide 9: Contact
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Contact & Acknowledgments"
    content.text = """Contact:
• Santosh Kumar
• santosh.kumar.mtech2022@sitpune.edu.in
• Symbiosis Institute of Technology
• GitHub: github.com/santoshray02/ExplainableCropAI

Acknowledgments:
• Symbiosis Institute of Technology
• Faculty guidance and mentorship
• Academic research community"""
    
    # Save
    prs.save("presentations/ExplainableCropAI_Modern_Clean.pptx")
    print("Presentation created successfully!")

if __name__ == "__main__":
    create_presentation()
